"""CR-002 parser, authorization, transaction and regression coverage."""
import io
import sys
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as api
import db
from resume_processing import process_resume, normalize_skill, UNSUPPORTED, UNREADABLE, TOO_LARGE, MAX_BYTES


def pdf_bytes(text='Python React AWS MySQL Rust'):
    """Produce a minimal valid text PDF without an extra test dependency."""
    stream = f'BT /F1 12 Tf 30 100 Td ({text}) Tj ET'.encode()
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>', b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
    data = b'%PDF-1.4\n'; offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += str(i).encode()+b' 0 obj\n'+obj+b'\nendobj\n'
    start = len(data)
    data += b'xref\n0 6\n0000000000 65535 f \n'
    data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    return data+f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()


class ResumeTests(unittest.TestCase):
    def test_pdf_and_normalization(self):
        self.assertEqual(process_resume('CV.PDF', pdf_bytes()), ['aws', 'mysql', 'python', 'react', 'rust'])
        self.assertEqual(normalize_skill('  Node.js  '), 'node.js')
        self.assertEqual(normalize_skill(' MACHINE   Learning '), 'machine learning')

    def test_docx_paragraphs_and_tables(self):
        from docx import Document
        doc = Document(); doc.add_paragraph('PYTHON React AWS')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'MySQL Rust'
        buffer = io.BytesIO(); doc.save(buffer)
        self.assertEqual(process_resume('CV.docx', buffer.getvalue()), ['aws', 'mysql', 'python', 'react', 'rust'])

    def test_rejects_unsupported_corrupt_empty_and_oversized(self):
        for name, body, error in [('cv.txt', b'Python', UNSUPPORTED), ('cv.pdf', b'broken', UNREADABLE), ('cv.docx', b'broken', UNREADABLE), ('cv.pdf', pdf_bytes(''), UNREADABLE), ('cv.pdf', b'x'*(MAX_BYTES+1), TOO_LARGE)]:
            with self.subTest(name=name, error=error):
                with self.assertRaisesRegex(ValueError, error): process_resume(name, body)

    def test_encrypted_pdf_rejected(self):
        with patch('pdfminer.pdfdocument.PDFDocument') as parser:
            parser.return_value.encryption = {'encrypted': True}
            with self.assertRaisesRegex(ValueError, UNREADABLE): process_resume('cv.pdf', pdf_bytes())


class APITests(unittest.TestCase):
    def setUp(self):
        self.secret = patch.object(api, 'JWT_SECRET', 'x'*40); self.secret.start()
        self.addCleanup(self.secret.stop)
        self.client = api.app.test_client()
        self.headers = {'Authorization': 'Bearer '+api.create_token(7)}

    def test_upload_and_exact_message(self):
        with patch.object(db, 'save_resume') as save, patch.object(db, 'flag_resume_skills', return_value=2) as flag:
            response = self.client.post('/api/resume/upload', headers=self.headers, data={'resume': (io.BytesIO(pdf_bytes()), 'cv.pdf')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['message'], 'Resume processed. 5 skills identified. 2 new skill(s) found missing from your profile. Choose Add to Profile below to include them.')
        save.assert_called_once()
        flag.assert_called_once_with(7, ['aws', 'mysql', 'python', 'react', 'rust'])

    def test_owner_and_admin_authorization(self):
        self.assertEqual(self.client.post('/api/resume/upload').status_code, 401)
        self.assertEqual(self.client.post('/api/students/8/resume', headers=self.headers).status_code, 403)
        with patch.object(db, 'get_user_by_id', return_value={'role': 'student'}):
            for path in ['/api/admin/approve-skill/1', '/api/admin/reject-skill/1']:
                self.assertEqual(self.client.post(path, headers=self.headers).status_code, 403)
            self.assertEqual(self.client.get('/api/admin/flagged-skills', headers=self.headers).status_code, 403)

    def test_admin_can_review_and_conflicts_are_reported(self):
        with patch.object(db, 'get_user_by_id', return_value={'role': 'admin'}), patch.object(db, 'review_flagged_skill', return_value=True) as review:
            self.assertEqual(self.client.post('/api/admin/approve-skill/3', headers=self.headers).status_code, 200)
            review.assert_called_once_with(3, 7, True)
            review.return_value = False
            self.assertEqual(self.client.post('/api/admin/reject-skill/3', headers=self.headers).status_code, 409)

    def test_get_with_json_content_type_and_no_body(self):
        with patch.object(db, 'get_jobs', return_value=[]):
            response = self.client.get('/api/jobs', headers={**self.headers, 'Content-Type': 'application/json'})
        self.assertEqual(response.status_code, 200)

    def test_remember_does_not_extend_24_hour_expiry(self):
        import jwt, time
        token = jwt.decode(api.create_token(7, remember=True), 'x'*40, algorithms=['HS256'])
        self.assertAlmostEqual(token['exp']-time.time(), 86400, delta=3)

    def test_malformed_json_and_large_upload(self):
        self.assertEqual(self.client.post('/api/auth/login', json=[]).status_code, 400)
        self.assertEqual(self.client.post('/api/auth/login', json={'email': 2}).status_code, 400)
        response = self.client.post('/api/resume/upload', headers=self.headers, data={'resume': (io.BytesIO(b'x'*(MAX_BYTES+100000)), 'cv.pdf')})
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json['error'], TOO_LARGE)


class DatabaseWorkflowTests(unittest.TestCase):
    def connection(self, results):
        conn = MagicMock(); cur = conn.cursor.return_value
        cur.fetchone.side_effect = results
        context = MagicMock(); context.__enter__.return_value = conn
        return conn, cur, patch.object(db, 'get_connection', return_value=context)

    def test_matched_duplicate_and_new_flag(self):
        conn, cur, context = self.connection([(7,), (1,), None, (4,), None, None])
        with context: self.assertEqual(db.flag_resume_skills(7, ['Python', 'RUST', '  python ', 'react']), 1)
        inserts = [call for call in cur.execute.call_args_list if call.args[0].startswith('INSERT INTO Unmatched')]
        self.assertEqual(len(inserts), 1); conn.commit.assert_called_once()
        lookups = [call for call in cur.execute.call_args_list if 'FROM User_Skills us' in call.args[0]]
        self.assertEqual(len(lookups), 3)
        self.assertTrue(all(call.args[1][0] == 7 for call in lookups))

    def test_approve_and_reject(self):
        for approve in [True, False]:
            conn, cur, context = self.connection([{'status': 'pending', 'skill_name': 'rust'}])
            with context: self.assertTrue(db.review_flagged_skill(1, 9, approve))
            inserts = [call for call in cur.execute.call_args_list if call.args[0].startswith('INSERT INTO Skills')]
            self.assertEqual(len(inserts), int(approve)); conn.commit.assert_called_once()

    def test_failure_rolls_back(self):
        conn, cur, context = self.connection([(7,)])
        cur.execute.side_effect = RuntimeError('database unavailable')
        with context, self.assertRaises(RuntimeError): db.flag_resume_skills(7, ['rust'])
        conn.rollback.assert_called_once(); conn.commit.assert_not_called()

    def test_skill_gap_regression(self):
        job = {'job_id': 1, 'skills': [
            {'skill_name': 'python', 'requirement_type': 'required'}, {'skill_name': 'sql', 'requirement_type': 'required'}, {'skill_name': 'aws', 'requirement_type': 'preferred'}]}
        conn, cur, context = self.connection([])
        cur.fetchall.return_value = [{'job_id': 1, 'skill_name': 'python'}, {'job_id': 1, 'skill_name': 'sql'}]
        with context, patch.object(db, 'get_saved_resume', return_value=None), patch.object(db, 'get_user_skills', return_value=['Python']), patch.object(db, 'get_job_by_id', return_value=job):
            result = db.get_skill_gap(1, 1)
            self.assertEqual({key: result[key] for key in ['matched', 'missing', 'preferred', 'match_pct', 'total_required']}, {'matched': ['python'], 'missing': ['sql'], 'preferred': ['aws'], 'match_pct': 50, 'total_required': 2})
            self.assertEqual(result['resume_match_source'], 'profile')



if __name__ == '__main__': unittest.main()
