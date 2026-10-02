"""App."""

from dotenv import load_dotenv
load_dotenv()

import os
import re
import secrets
import smtplib
from email.message import EmailMessage
import jwt
import bcrypt
import datetime
from functools import wraps
from flask import Flask, request, jsonify, send_file
from werkzeug.utils import secure_filename
from flask_cors import CORS

import db
from resume_processing import process_resume, MAX_BYTES, TOO_LARGE
from werkzeug.exceptions import HTTPException
from io import BytesIO
from flask import Request

class MemoryUploadRequest(Request):
    """Keep bounded multipart uploads in memory instead of temporary files."""
    def _get_file_stream(self, *args, **kwargs):
        return BytesIO()


app = Flask(__name__)
app.request_class = MemoryUploadRequest
origins = os.environ.get("CORS_ORIGINS", "").split(",")
if any(origins):
    CORS(app, origins=[origin.strip() for origin in origins if origin.strip()])
app.config["MAX_CONTENT_LENGTH"] = MAX_BYTES + 64 * 1024
ALLOWED_RESUME_EXTENSIONS = {"pdf", "docx"}

JWT_SECRET  = os.environ.get("JWT_SECRET")
JWT_EXPIRES = 24  # hours
EMAIL_PATTERN = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
PASSWORD_PATTERN = re.compile(r"^(?=.*[a-z])(?=.*[A-Z])(?=.*\d).{8,}$")


# JWT helper

def create_token(user_id, remember=False):
    if not JWT_SECRET:
        raise RuntimeError("JWT_SECRET must be configured before authentication can be used")
    payload = {
        "user_id": user_id,
        "exp": datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=JWT_EXPIRES),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm="HS256")


def token_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        auth = request.headers.get("Authorization", "")
        if not auth.startswith("Bearer "):
            return jsonify({"error": "Missing or invalid token"}), 401
        token = auth.split(" ", 1)[1]
        try:
            payload = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
            request.user_id = payload["user_id"]
        except jwt.ExpiredSignatureError:
            return jsonify({"error": "Token expired"}), 401
        except jwt.InvalidTokenError:
            return jsonify({"error": "Invalid token"}), 401
        return f(*args, **kwargs)
    return decorated


def _serialize(obj):
    """Convert non-JSON-serializable types (e.g. date) to strings."""
    if isinstance(obj, (datetime.date, datetime.datetime)):
        return obj.isoformat()
    return str(obj)


def jsonify_safe(data):
    import json
    return app.response_class(
        response=json.dumps(data, default=_serialize),
        status=200,
        mimetype="application/json",
    )


def password_error(password):
    """Return an actionable password-policy message, or None when valid."""
    if not isinstance(password, str) or len(password.encode("utf-8")) > 72:
        return "Password must be at most 72 UTF-8 bytes"
    if not PASSWORD_PATTERN.fullmatch(password or ""):
        return "Password must be 8+ characters and include uppercase, lowercase, and a number"
    return None


def allowed_resume(filename):
    """Accept only the two resume formats required by the SRS."""
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_RESUME_EXTENSIONS


def deliver_password_reset(email, token):
    """Send the reset URL using deployment-provided SMTP credentials."""
    host = os.environ.get("SMTP_HOST")
    sender = os.environ.get("SMTP_FROM")
    if not host or not sender:
        app.logger.warning("Password reset requested but SMTP is not configured")
        return False
    public_url = os.environ.get("APP_PUBLIC_URL", "http://127.0.0.1:3000").rstrip("/")
    message = EmailMessage()
    message["Subject"] = "NovaTeck password reset"
    message["From"], message["To"] = sender, email
    message.set_content(f"Reset your NovaTeck password within 30 minutes: {public_url}/?resetToken={token}")
    try:
        with smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", "587")), timeout=15) as server:
            if os.environ.get("SMTP_STARTTLS", "true").lower() == "true":
                server.starttls()
            username, password = os.environ.get("SMTP_USER"), os.environ.get("SMTP_PASSWORD")
            if username and password:
                server.login(username, password)
            server.send_message(message)
        return True
    except (OSError, smtplib.SMTPException) as exc:
        app.logger.error("Password reset email failed: %s", exc)
        return False


# AUTH

@app.route("/api/auth/register", methods=["POST"])
def register():
    data     = request.get_json() or {}
    name     = data.get("name", "").strip()
    email    = data.get("email", "").strip().lower()
    password = data.get("password", "")
    skills   = data.get("skills", "")  # comma-separated string or list

    if not name or not email or not password:
        return jsonify({"error": "name, email, and password are required"}), 400
    if not EMAIL_PATTERN.fullmatch(email):
        return jsonify({"error": "Enter a valid email address"}), 400
    if error := password_error(password):
        return jsonify({"error": error}), 400

    pw_hash = bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()

    # Normalize skills to comma-separated string
    if isinstance(skills, list):
        skills_str = ",".join(skills)
    else:
        skills_str = skills

    user_id = db.insert_user(name, email, pw_hash, skills_str)
    if not user_id:
        return jsonify({"error": "Email already registered"}), 409

    # Populate User_Skills table from skills string
    if skills_str:
        skill_names = [s.strip() for s in skills_str.split(",") if s.strip()]
        db.set_user_skills_from_list(user_id, skill_names)

    token = create_token(user_id)
    return jsonify({"token": token, "user_id": user_id, "name": name}), 201


@app.route("/api/auth/login", methods=["POST"])
def login():
    data     = request.get_json() or {}
    email    = data.get("email", "").strip().lower()
    password = data.get("password", "")
    remember = bool(data.get("remember"))

    if not email or not password:
        return jsonify({"error": "email and password are required"}), 400
    if not EMAIL_PATTERN.fullmatch(email):
        return jsonify({"error": "Enter a valid email address"}), 400

    if not isinstance(password, str) or len(password.encode("utf-8")) > 72:
        return jsonify({"error": "Invalid email or password"}), 401
    user = db.get_user_by_email(email)
    if not user:
        return jsonify({"error": "Invalid email or password"}), 401

    if not bcrypt.checkpw(password.encode(), user["password_hash"].encode()):
        return jsonify({"error": "Invalid email or password"}), 401

    token = create_token(user["user_id"], remember=remember)
    return jsonify({
        "token":   token,
        "user_id": user["user_id"],
        "name":    user["name"],
        "role":    user.get("role", "student"),
    }), 200


@app.route("/api/auth/password-reset/request", methods=["POST"])
def request_password_reset():
    """Request password reset."""
    email = (request.get_json() or {}).get("email", "").strip().lower()
    if not EMAIL_PATTERN.fullmatch(email):
        return jsonify({"error": "Enter a valid email address"}), 400
    if not os.environ.get('SMTP_HOST') or not os.environ.get('SMTP_FROM'):
        return jsonify({'error': 'Password reset email is not available yet. Please contact support.'}), 503
    user = db.get_user_by_email(email)
    if user:
        token = secrets.token_urlsafe(32)
        if not db.create_password_reset_token(user["user_id"], token) or not deliver_password_reset(email, token):
            return jsonify({'error': 'Could not send the reset email. Please try again later or contact support.'}), 503
    return jsonify({"message": "If the email exists, a password-reset link has been sent."}), 200


@app.route("/api/auth/password-reset/confirm", methods=["POST"])
def confirm_password_reset():
    data = request.get_json() or {}
    token, password = data.get("token", ""), data.get("password", "")
    if error := password_error(password):
        return jsonify({"error": error}), 400
    password_hash = bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()
    if not db.reset_password(token, password_hash):
        return jsonify({"error": "That reset link is invalid or has expired"}), 400
    return jsonify({"message": "Password updated. You can now sign in."}), 200


# JOBS

@app.route("/api/jobs", methods=["GET"])
@token_required
def get_jobs():
    skill      = request.args.getlist("skill")
    location   = request.args.get("location")
    company_id = request.args.get("company_id", type=int)
    experience = request.args.get("experience", type=str)
    discipline = request.args.getlist("discipline")
    limit      = request.args.get("limit", 200, type=int)
    offset     = request.args.get("offset", 0, type=int)
    limit = max(1, min(limit or 200, 500))
    offset = max(0, offset or 0)

    jobs = db.get_jobs(
        skill=skill,
        location=location,
        company_id=company_id,
        experience=experience,
        discipline=discipline,
        limit=None if request.args.get("sort") == "match" else limit,
        offset=offset,
        q=request.args.get("q"),
        work_type=request.args.get("work_type"),
        exclude_applications_user_id=request.user_id,
    )
    jobs = db.add_resume_matches(request.user_id, jobs)
    if request.args.get('sort') == 'match':
        jobs.sort(key=lambda job: job.get('resume_match_pct') if isinstance(job.get('resume_match_pct'), (int, float)) else -1, reverse=True)
        jobs = jobs[offset:offset + limit]
    return jsonify_safe(jobs)


@app.route("/api/jobs/count", methods=["GET"])
@token_required
def get_job_count():
    """Return a filtered total for pagination without transferring every row."""
    return jsonify({"total": db.get_job_count(
        q=request.args.get("q"),
        work_type=request.args.get("work_type"),
        exclude_applications_user_id=request.user_id,
        skill=request.args.getlist("skill"),
        location=request.args.get("location"),
        company_id=request.args.get("company_id", type=int),
        experience=request.args.get("experience", type=str),
        discipline=request.args.getlist("discipline"),
    )})


@app.route("/api/search-options", methods=["GET"])
@token_required
def search_options():
    return jsonify_safe(db.get_search_options())


@app.route("/api/jobs/<int:job_id>", methods=["GET"])
@token_required
def get_job(job_id):
    job = db.get_job_by_id(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    return jsonify_safe(job)


# SAVED JOBS

@app.route("/api/saved-jobs", methods=["POST"])
@token_required
def save_job():
    data   = request.get_json() or {}
    job_id = data.get("job_id")
    if not job_id:
        return jsonify({"error": "job_id required"}), 400

    result = db.save_job(request.user_id, job_id)
    if result:
        return jsonify({"message": "Job saved"}), 201
    return jsonify({"message": "Already saved"}), 200


@app.route("/api/saved-jobs/<int:job_id>", methods=["DELETE"])
@token_required
def remove_saved_job(job_id):
    db.remove_saved_job(request.user_id, job_id)
    return jsonify({"message": "Removed"}), 200


@app.route("/api/students/<int:user_id>/saved-jobs", methods=["GET"])
@token_required
def get_saved_jobs(user_id):
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    jobs = db.get_saved_jobs(user_id)
    return jsonify_safe(db.add_resume_matches(request.user_id, jobs))


# STUDENT PROFILE & SKILLS

@app.route("/api/students/<int:user_id>", methods=["GET"])
@token_required
def get_student(user_id):
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    user = db.get_user_by_id(user_id)
    if not user:
        return jsonify({"error": "User not found"}), 404
    user.pop("password_hash", None)
    user.pop("resume_url", None)
    user.pop("resume_filename", None)
    return jsonify_safe(user)


@app.route("/api/students/<int:user_id>", methods=["PUT"])
@token_required
def update_student(user_id):
    """Update the editable profile fields without exposing password data."""
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    name = (request.get_json() or {}).get("name", "").strip()
    if not name:
        return jsonify({"error": "Name is required"}), 400
    if not db.update_user(user_id, name=name):
        return jsonify({"error": "Profile was not updated"}), 400
    return jsonify({"message": "Profile updated", "name": name})


@app.route("/api/resume/upload", methods=["POST"])
@app.route("/api/students/<int:user_id>/resume", methods=["POST"])
@token_required
def upload_resume(user_id=None):
    """Upload resume."""
    if user_id is not None and request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    resume = request.files.get("resume")
    if not resume or not resume.filename:
        return jsonify({"error": "Choose a PDF or DOCX resume"}), 400
    try:
        content = resume.read(MAX_BYTES + 1)
        skills = process_resume(resume.filename, content)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    db.save_resume(request.user_id, secure_filename(resume.filename)[-255:] or "resume.pdf", content, skills)
    count = db.flag_resume_skills(request.user_id, skills)
    return jsonify({"message": f"Resume processed. {len(skills)} skills identified. {count} new skill(s) found missing from your profile. Choose Add to Profile below to include them.",
                    "skills": skills, "flagged_count": count}), 200


@app.route("/api/resume/saved", methods=["GET"])
@token_required
def saved_resume_metadata():
    return jsonify_safe(db.get_saved_resume(request.user_id))


@app.route("/api/resume/saved/file", methods=["GET"])
@token_required
def saved_resume_file():
    resume = db.get_saved_resume(request.user_id, include_content=True)
    if not resume:
        return jsonify({"error": "No saved resume found"}), 404
    is_pdf = resume['filename'].lower().endswith('.pdf')
    response = send_file(BytesIO(resume['content']), download_name=resume['filename'],
        mimetype='application/pdf' if is_pdf else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=not is_pdf)
    response.headers['Content-Security-Policy'] = "sandbox"
    return response


@app.route("/api/resume/saved/preview", methods=["GET"])
@token_required
def saved_resume_preview():
    resume = db.get_saved_resume(request.user_id, include_content=True)
    if not resume:
        return jsonify({"error": "No saved resume found"}), 404
    if not resume['filename'].lower().endswith('.docx'):
        return jsonify({"error": "Text preview supports DOCX resumes"}), 400
    try:
        from docx import Document
        document = Document(BytesIO(resume['content']))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(' | '.join(cell.text for cell in row.cells) for row in table.rows)
        response = jsonify({"text": '\n'.join(paragraphs)})
        response.headers['Cache-Control'] = 'no-store'
        return response
    except Exception:
        return jsonify({"error": "Could not preview this resume"}), 422


@app.route("/api/resume/saved/scan", methods=["POST"])
@token_required
def rescan_saved_resume():
    resume = db.get_saved_resume(request.user_id, include_content=True)
    if not resume:
        return jsonify({"error": "Upload a resume first"}), 404
    try:
        skills = process_resume(resume['filename'], resume['content'])
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    count = db.flag_resume_skills(request.user_id, skills)
    return jsonify({"skills": skills, "flagged_count": count,
                    "message": f"Saved resume scanned. {len(skills)} skills identified. {count} new skills to add."})


def admin_required(f):
    """Check the current database role rather than trusting client claims."""
    @wraps(f)
    @token_required
    def decorated(*args, **kwargs):
        user = db.get_user_by_id(request.user_id)
        if not user or user.get("role") != "admin":
            return jsonify({"error": "Administrator access required"}), 403
        return f(*args, **kwargs)
    return decorated


@app.route("/api/resume/skills/<int:flag_id>/add", methods=["POST"])
@token_required
def add_detected_skill(flag_id):
    if not db.add_resume_skill_to_profile(request.user_id, flag_id):
        return jsonify({"error": "Detected skill not found for your account"}), 404
    return jsonify({"message": "Skill added to your profile"})


@app.route('/api/resume/skills/<int:flag_id>', methods=['DELETE'])
@token_required
def dismiss_detected_skill(flag_id):
    if not db.dismiss_resume_skill(request.user_id, flag_id):
        return jsonify({'error': 'Pending skill not found for your account.'}), 404
    return jsonify({'message': 'Flagged skill dismissed.'})


@app.route("/api/admin/flagged-skills")
@admin_required
def flagged_skills():
    """List pending skills for authorized administrators."""
    return jsonify_safe(db.get_flagged_skills())


@app.route("/api/admin/approve-skill/<int:flag_id>", methods=["POST"])
@admin_required
def approve_skill(flag_id):
    """Approve an existing pending flag."""
    if not db.review_flagged_skill(flag_id, request.user_id, True):
        return jsonify({"error": "Flag not found or already reviewed"}), 409
    return jsonify({"message": "Skill approved"})


@app.route("/api/admin/reject-skill/<int:flag_id>", methods=["POST"])
@admin_required
def reject_skill(flag_id):
    """Reject a flag without changing the skill catalog."""
    if not db.review_flagged_skill(flag_id, request.user_id, False):
        return jsonify({"error": "Flag not found or already reviewed"}), 409
    return jsonify({"message": "Skill rejected"})


@app.route("/api/students/<int:user_id>/skills", methods=["PUT"])
@token_required
def update_skills(user_id):
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    data   = request.get_json() or {}
    skills = data.get("skills", [])
    if isinstance(skills, str):
        skills = [s.strip() for s in skills.split(",") if s.strip()]

    if not isinstance(skills, list) or any(not isinstance(s, str) or len(s) > 100 for s in skills) or len(skills) > 100:
        return jsonify({"error": "Provide at most 100 skill names, each at most 100 characters"}), 400
    db.set_user_skills_from_list(user_id, skills)
    return jsonify({"message": "Skills updated", "skills": skills}), 200


# SKILL GAP ANALYSIS

@app.route("/api/students/<int:user_id>/skill-gap", methods=["GET"])
@token_required
def skill_gap(user_id):
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403

    job_id = request.args.get("job_id", type=int)

    if job_id:
        # Single-job skill gap
        result = db.get_skill_gap(user_id, job_id)
    else:
        # Full market skill gap summary
        result = db.get_student_skill_gap_summary(user_id, request.args.get("discipline", type=str))

    result["unverified_skills"] = db.get_flagged_skills(user_id)
    return jsonify_safe(result)


# COMPANIES

@app.route("/api/companies", methods=["GET"])
@token_required
def get_companies():
    companies = db.get_all_companies()
    return jsonify_safe(companies)


@app.route("/api/jobs/sync-status", methods=["GET"])
@token_required
def job_sync_status():
    """Job sync status."""
    run = db.get_latest_sync_run()
    return jsonify_safe(run or {"status": "not_configured"})


@app.route("/api/students/<int:user_id>/recommendations")
@token_required
def recommendations(user_id):
    """Return active jobs ranked by the student's verified catalog skills."""
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    return jsonify_safe(db.add_resume_matches(user_id, db.get_recommendations(user_id, db.get_profile_details(user_id).get('career_focus'))))


@app.route("/api/market-trends")
@token_required
def market_trends():
    """Expose the recorded historical demand series."""
    return jsonify_safe(db.get_market_trends())


@app.route("/api/students/<int:user_id>/profile-details", methods=["GET", "PUT"])
@token_required
def profile_details(user_id):
    """Read or save private profile details for the authenticated owner."""
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    if request.method == "GET":
        return jsonify_safe(db.get_profile_details(user_id))
    from profile_validation import validate_details
    data = request.get_json() or {}
    name, email = data.get("name", "").strip(), data.get("email", "").strip().lower()
    if not name or len(name) > 255 or not EMAIL_PATTERN.fullmatch(email):
        return jsonify({"error": "Enter a full name and valid email address"}), 400
    try:
        details = validate_details(data.get("details", {}))
        existing = db.get_user_by_email(email)
        if existing and existing["user_id"] != user_id:
            return jsonify({"error": "Email already registered"}), 409
        db.save_profile_details(user_id, name, email, details)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    return jsonify({"message": "Profile updated", "name": name})


@app.route("/api/students/<int:user_id>/applications", methods=["GET", "POST"])
@token_required
def application_activity(user_id):
    """Track owner-reported application activity."""
    if request.user_id != user_id:
        return jsonify({"error": "Unauthorized"}), 403
    if request.method == "GET":
        return jsonify_safe(db.add_resume_matches(user_id, db.get_application_activity(user_id)))
    data = request.get_json() or {}
    job_id, status = data.get("job_id"), data.get("status", "Opened employer site")
    if not isinstance(job_id, int) or status not in ['Opened employer site', 'Applied', 'Interviewing', 'Offer', 'Closed']:
        return jsonify({"error": "Invalid application details"}), 400
    if not db.get_job_by_id(job_id):
        return jsonify({"error": "Job not found"}), 404
    db.set_application_activity(user_id, job_id, status)
    return jsonify({"message": "Application activity updated"})


@app.route("/api/admin/overview")
@admin_required
def admin_overview():
    """Return authenticated administrator dashboard data."""
    return jsonify_safe(db.get_admin_overview())


# HEALTH CHECK

@app.route("/api/health", methods=["GET"])
def health():
    database_ok = db.database_is_healthy()
    return jsonify({
        "status": "ok" if database_ok else "degraded",
        "project": "NovaTeck DFW Tech Job Tracker",
        "database": "ok" if database_ok else "unavailable",
    }), 200 if database_ok else 503


@app.before_request
def validate_json_object():
    """Reject unexpected JSON shapes before route-specific processing."""
    if request.method in {"POST", "PUT", "PATCH"} and request.is_json:
        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            return jsonify({"error": "A JSON object is required"}), 400
        for field in ("name", "email", "password", "token"):
            if field in data and (not isinstance(data[field], str) or len(data[field]) > 1000):
                return jsonify({"error": f"Invalid {field}"}), 400
        skills = data.get("skills", [])
        if not isinstance(skills, (str, list)) or (isinstance(skills, list) and any(not isinstance(s, str) or len(s) > 100 for s in skills)):
            return jsonify({"error": "Invalid skills"}), 400


@app.after_request
def security_headers(response):
    """Prevent sniffing, framing and caching of private API responses."""
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Cache-Control"] = "no-store"
    return response


@app.errorhandler(413)
def upload_too_large(error):
    return jsonify({"error": TOO_LARGE}), 413


@app.errorhandler(Exception)
def api_error(error):
    """Return JSON errors without leaking internal exception details."""
    if isinstance(error, HTTPException):
        return jsonify({"error": error.description}), error.code
    app.logger.exception("Request failed")
    return jsonify({"error": "Service temporarily unavailable. Please try again."}), 503


@app.route('/api/support', methods=['POST'])
@token_required
def submit_support():
    body = request.get_json(silent=True)
    if not isinstance(body, dict):
        return jsonify({'error': 'Enter a subject and message.'}), 400
    subject, message = body.get('subject'), body.get('message')
    if not isinstance(subject, str) or not 1 <= len(subject.strip()) <= 150 or not isinstance(message, str) or not 1 <= len(message.strip()) <= 3000:
        return jsonify({'error': 'Use a subject of 1–150 characters and a message of 1–3,000 characters.'}), 400
    return jsonify({'support_id': db.create_support_request(request.user_id, subject.strip(), message.strip())}), 201


@app.route('/api/admin/support', methods=['GET'])
@admin_required
def support_inbox():
    return jsonify_safe(db.support_threads())


@app.route('/api/support', methods=['GET'])
@token_required
def my_support_requests():
    return jsonify_safe(db.support_threads(request.user_id))


@app.route('/api/admin/support/<int:support_id>/replies', methods=['POST'])
@admin_required
def support_reply(support_id):
    body = request.get_json(silent=True)
    message = body.get('message') if isinstance(body, dict) else None
    if not isinstance(message, str) or not 1 <= len(message.strip()) <= 3000:
        return jsonify({'error': 'Enter a reply of 1–3,000 characters.'}), 400
    if not db.reply_to_support(support_id, request.user_id, message.strip()):
        return jsonify({'error': 'Support request not found.'}), 404
    return jsonify({'message': 'Reply sent.'}), 201


@app.route('/api/admin/support/<int:support_id>', methods=['PATCH'])
@admin_required
def support_status(support_id):
    body = request.get_json(silent=True)
    status = body.get('status') if isinstance(body, dict) else None
    if status not in ('Open', 'Resolved'):
        return jsonify({'error': 'Choose Open or Resolved.'}), 400
    if not db.update_support_request(support_id, status):
        return jsonify({'error': 'Support request not found.'}), 404
    return jsonify({'status': status})


@app.route('/api/chat', methods=['POST'])
@token_required
def career_chat():
    from career_chat import validate_messages, reply
    data = request.get_json() or {}
    try:
        messages = validate_messages(data.get('messages') if isinstance(data, dict) else None)
    except ValueError as error:
        return jsonify({'error': str(error)}), 400
    result, status = reply(messages)
    return jsonify(result), status




@app.route('/api/students/<int:user_id>/applications/<int:job_id>/notes', methods=['PUT'])
@token_required
def application_notes(user_id, job_id):
    if request.user_id != user_id:
        return jsonify({'error': 'Unauthorized'}), 403
    data = request.get_json() or {}
    notes, interview_date = data.get('notes', ''), data.get('interview_date') or None
    if not isinstance(notes, str) or len(notes) > 3000:
        return jsonify({'error': 'Notes must be at most 3000 characters'}), 400
    if interview_date:
        try:
            from datetime import date
            date.fromisoformat(interview_date)
        except (ValueError, TypeError):
            return jsonify({'error': 'Enter a valid interview date'}), 400
    if not db.save_application_notes(user_id, job_id, notes.strip(), interview_date):
        return jsonify({'error': 'Application not found'}), 404
    return jsonify({'message': 'Application notes saved'})


@app.route('/api/hidden-jobs/<int:job_id>', methods=['PUT', 'DELETE'])
@token_required
def hidden_job(job_id):
    if not db.get_job_by_id(job_id):
        return jsonify({'error': 'Job not found'}), 404
    db.set_hidden_job(request.user_id, job_id, request.method == 'PUT')
    return jsonify({'message': 'Job hidden' if request.method == 'PUT' else 'Job restored'})


@app.route('/api/hidden-jobs/restore', methods=['POST'])
@token_required
def restore_hidden_jobs():
    count = db.restore_hidden_jobs(request.user_id)
    return jsonify({'message': f'{count} hidden jobs restored', 'count': count})


if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1", port=5000)
