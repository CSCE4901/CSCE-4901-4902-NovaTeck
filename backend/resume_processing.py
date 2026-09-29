"""Resume processing."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from nlp_tagger import tag_skills_for_job

MAX_BYTES = 5 * 1024 * 1024
UNSUPPORTED = "Unsupported file type. Please upload a PDF or DOCX file."
UNREADABLE = "Resume could not be processed. Please upload a text-based PDF or DOCX file."
TOO_LARGE = "File size exceeds the 5 MB limit. Please upload a smaller file."


def normalize_skill(name):
    """Canonicalize case and whitespace before exact comparison."""
    return " ".join(name.lower().split())


def process_resume(filename, content):
    """Extract text and skills in memory; never persist a resume or its text."""
    extension = filename.rsplit('.', 1)[-1].lower()
    if extension not in {'pdf', 'docx'}:
        raise ValueError(UNSUPPORTED)
    if len(content) > MAX_BYTES:
        raise ValueError(TOO_LARGE)
    try:
        if extension == 'pdf':
            from pdfminer.high_level import extract_text
            from pdfminer.pdfparser import PDFParser
            from pdfminer.pdfdocument import PDFDocument
            document = PDFDocument(PDFParser(BytesIO(content)))
            if document.encryption:
                raise ValueError(UNREADABLE)
            text = extract_text(BytesIO(content))
        else:
            from docx import Document
            with ZipFile(BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 25 * MAX_BYTES:
                    raise ValueError(UNREADABLE)
            doc = Document(BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
            text = '\n'.join(paragraphs)
        if not text.strip():
            raise ValueError(UNREADABLE)
    except Exception as exc:
        raise ValueError(UNREADABLE) from exc
    skills = sorted({normalize_skill(item['skill_name']) for item in tag_skills_for_job(text)})
    return skills
